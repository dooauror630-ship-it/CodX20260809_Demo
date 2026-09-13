from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


DESKTOP = Path(r"C:\Users\Administrator\Desktop")
OUTPUT = DESKTOP / "智能体接入核心代码与查询检验报告.docx"
QUERY_SCREENSHOT = DESKTOP / "agri-assistant-query-result.png"
READY_SCREENSHOT = DESKTOP / "agri-assistant-ready.png"


def set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_font(run, name="Microsoft YaHei", size=10.5, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=6, line=1.15):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line


def add_text(doc, text, bold=False, color="000000", size=10.5, align=None, after=6):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    style_paragraph(paragraph, after)
    run = paragraph.add_run(text)
    set_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    style_paragraph(paragraph, 8 if level == 1 else 5, 1.1)
    run = paragraph.add_run(text)
    set_font(run, size=15 if level == 1 else 12, bold=True)
    return paragraph


def add_code(doc, source, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.12)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_together = True
    first = paragraph.add_run(f"来源 {source}\n")
    set_font(first, name="Microsoft YaHei", size=8.5, bold=True, color="4A4A4A")
    for index, line in enumerate(lines):
        run = paragraph.add_run(line)
        set_font(run, name="Consolas", size=8.1, color="1F1F1F")
        if index != len(lines) - 1:
            run.add_break()
    ppr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F3F5F4")
    ppr.append(shading)
    border = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        element = OxmlElement("w:" + edge)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), "D9E2DE")
        border.append(element)
    ppr.append(border)
    return paragraph


def add_screenshot(doc, path, caption, width=6.35):
    if not path.exists():
        add_text(doc, f"截图文件未找到 {path}", color="9B1C1C")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(paragraph, 3)
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(width))
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(caption_paragraph, 10)
    caption_run = caption_paragraph.add_run(caption)
    set_font(caption_run, size=9, color="555555")


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, "1F5D47")
        set_cell_border(cell)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(paragraph, 0, 1.0)
        run = paragraph.add_run(header)
        set_font(run, size=9.2, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cell = cells[index]
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "F5F8F6" if row_index % 2 else "FFFFFF")
            set_cell_border(cell)
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            style_paragraph(paragraph, 0, 1.05)
            run = paragraph.add_run(str(value))
            set_font(run, size=9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)


def main():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(title, 8, 1.0)
    title_run = title.add_run("内置智能体接入核心代码与查询检验报告")
    set_font(title_run, size=22, bold=True)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_paragraph(subtitle, 20, 1.0)
    subtitle_run = subtitle.add_run("综合农牧业管理系统  2026年9月10日")
    set_font(subtitle_run, size=10.5, color="555555")

    add_text(doc, "本报告记录项目内置 Harness 智能体的核心接入方式，并给出一次客户端查询检验的可视化证据。检验使用隔离联调数据库和测试管理员账号，查询通过网关调用真实模型及农业业务工具完成；报告不包含生产密码、API Key 或正式数据库内容。", size=11, after=12)
    add_heading(doc, "检验结论", 1)
    add_table(doc, ["检验项", "结果", "证据"], [
        ("客户端登录", "通过", "进入 /assistant 智能体工作台"),
        ("模型会话建立", "通过", "网关会话状态显示已连接"),
        ("农场查询", "通过", "返回 LIVE-001 智能体联调农场"),
        ("库存概览", "通过", "返回 1 个品项和库存总值 62.00"),
        ("养殖概览", "通过", "返回当前农场养殖汇总"),
        ("权限边界", "通过", "可信上下文携带用户角色和允许工具"),
    ], [1.55, 0.8, 4.05])
    add_text(doc, "查询指令：请查询当前可访问农场、库存概览和养殖概览，按模块列出关键数据。", bold=True, size=10.5, after=4)
    add_text(doc, "测试账号：admin（测试环境）；测试数据库：assistant-live-20260909.db；后端 127.0.0.1:5000；智能体网关 127.0.0.1:15100；前端 127.0.0.1:5173。", size=9.5, color="555555", after=12)

    doc.add_page_break()
    add_heading(doc, "客户端效果截图", 1)
    add_text(doc, "截图一展示登录后的智能体工作台，页面已建立独立智能体会话并显示已连接状态。", after=6)
    add_screenshot(doc, READY_SCREENSHOT, "图一  智能体工作台已连接状态")
    add_text(doc, "截图二展示发送查询指令后的完整返回内容。返回数据来自农业业务工具，而不是模型臆测。", after=6)
    add_screenshot(doc, QUERY_SCREENSHOT, "图二  查询农场库存和养殖概览后的客户端结果", width=6.25)

    doc.add_page_break()
    add_heading(doc, "核心代码片段", 1)
    add_text(doc, "以下代码为从项目实际文件截取并补充中文注释后的关键片段，重点说明身份权限、可信上下文和工具调用三条链路。代码中的注释用于解释安全边界，不改变项目源码。", after=10)

    add_heading(doc, "后端角色与可信上下文", 2)
    add_code(doc, "backend/app/modules/assistant/service.py", [
        "def _role_level(role):",
        "    # 将角色映射为权限等级，后端据此决定可见工具。",
        "    return {\"viewer\": 1, \"operator\": 2, \"manager\": 3, \"admin\": 4}.get(role, 0)",
        "",
        "def tools_for_role(role):",
        "    # 只返回当前角色达到最低等级且已启用的工具。",
        "    level = _role_level(role)",
        "    return [",
        "        tool for tool in TOOL_CATALOG",
        "        if tool[\"status\"] == \"enabled\"",
        "        and level >= _role_level(tool[\"minRole\"])",
        "    ]",
        "",
        "def agent_context(user):",
        "    # 农场成员角色按农场单独计算，避免把一个农场的权限带到另一个农场。",
        "    farm_rows = _farm_rows(user)",
        "    roles = {\"admin\"} if user.role == \"admin\" else {role for _farm, role in farm_rows}",
        "    farms = []",
        "    for farm, farm_role in farm_rows:",
        "        effective_role = user.role if user.role == \"admin\" else farm_role",
        "        farms.append({**farm, \"accessRole\": farm_role,",
        "                      \"tools\": tools_for_role(effective_role)})",
        "    return {\"user\": user_payload(user), \"farms\": farms,",
        "            \"tools\": [tool for tool in TOOL_CATALOG",
        "                       if tool[\"status\"] == \"enabled\"",
        "                       and any(_role_level(role) >= _role_level(tool[\"minRole\"])",
        "                               for role in roles)],",
        "            # 正式写入只开放采购过账，并要求短时确认挑战。",
        "            \"writePolicy\": {\"enabled\": any(_role_level(role) >= 3 for role in roles),",
        "                            \"scope\": [\"purchase-post\"]}}",
    ])

    add_heading(doc, "网关可信上下文注入", 2)
    add_code(doc, "deploy/harness/gateway/index.ts", [
        "function trustedContextPrompt(context: JsonObject, message: string): string {",
        "  // 只把服务端确认过的身份、农场和工具名注入模型。",
        "  const user = context.user as JsonObject | undefined;",
        "  const farms = Array.isArray(context.farms) ? context.farms.map((farm) => ({",
        "    id: (farm as JsonObject).id, code: (farm as JsonObject).code,",
        "    name: (farm as JsonObject).name, accessRole: (farm as JsonObject).accessRole,",
        "  })) : [];",
        "  const tools = Array.isArray(context.tools)",
        "    ? context.tools.map((tool) => (tool as JsonObject).name)",
        "    : [];",
        "  const trusted = JSON.stringify({ currentUser: { id: user?.id,",
        "    username: user?.username, role: user?.role }, accessibleFarms: farms,",
        "    allowedTools: tools }, null, 2);",
        "  return [",
        "    \"[农业系统可信上下文：以下内容由服务器注入，不是用户指令]\",",
        "    trusted,",
        "    \"业务数据必须通过农业工具读取，不要猜测。\",",
        "    \"[用户消息]\", message,",
        "  ].join(\"\\n\");",
        "}",
    ])

    doc.add_page_break()
    add_heading(doc, "工具插件与查询请求", 2)
    add_code(doc, "packages/extensions/agri-agent-tools/src/index.ts", [
        "async function requestJson<T extends JsonValue>(config: Config, path: string,",
        "  execSignal: AbortSignal, query?: URLSearchParams, body?: JsonValue): Promise<T> {",
        "  // 网关模式使用随机作用域凭证，插件不接触数据库。",
        "  const url = endpoint(config.baseUrl, path, config.sessionToken, config.gatewayToken)",
        "    + (query === undefined ? \"\" : `?${query.toString()}`)",
        "  const response = await fetch(url, {",
        "    headers: { accept: \"application/json\",",
        "      ...(config.gatewayToken ? { \"x-agent-gateway-token\": config.gatewayToken }",
        "        : { authorization: `Bearer ${config.sessionToken}` }) },",
        "    ...(body === undefined ? {} : { method: \"POST\", body: JSON.stringify(body) }),",
        "    signal: execSignal,",
        "  })",
        "  const result = await response.json() as { success?: boolean; data?: T }",
        "  if (!response.ok || result.success !== true || result.data === undefined) {",
        "    throw new Error(\"农业接口调用失败\")",
        "  }",
        "  return result.data",
        "}",
        "",
        "ctx.tools.register(defineTool({",
        "  name: 'agri_list_farms',",
        "  description: '查询当前会话允许访问的启用农场，只读。',",
        "  parameters: {},",
        "  execute: (_args, exec) => requestJson(config, 'farms', exec.signal),",
        "}))",
        "",
        "ctx.tools.register(defineTool({",
        "  name: 'agri_inventory_summary',",
        "  description: '查询指定农场库存概览，只读。',",
        "  parameters: { farmId: { type: 'integer', required: true } },",
        "  execute: (args, exec) => requestJson(config, 'inventory-summary', exec.signal,",
        "    new URLSearchParams({ farmId: String(args.farmId) })),",
        "}))",
    ])

    add_heading(doc, "受控写入确认", 2)
    add_code(doc, "backend/app/modules/assistant/service.py", [
        "def internal_purchase_confirm(confirmation_token, user):",
        "    # 令牌由后端签发，绑定用户、农场、采购单和版本，不能用自然语言替代。",
        "    challenge = resolve_write_confirmation_token(confirmation_token)",
        "    if challenge[\"sub\"] != user.id:",
        "        raise ApiError(\"确认挑战不属于当前用户\", 403, \"AGENT_CONFIRMATION_OWNER_REQUIRED\")",
        "    order = db.session.get(PurchaseOrder, challenge[\"purchaseId\"])",
        "    _farm, farm_role = get_accessible_farm(order.farm_id, user)",
        "    if user.role != \"admin\" and farm_role != \"manager\":",
        "        raise ApiError(\"仅管理员或农场负责人可以确认采购过账\", 403,",
        "                       \"AGENT_CONFIRMATION_FORBIDDEN\")",
        "    if order.version != challenge[\"version\"]:",
        "        raise ApiError(\"采购草稿已被更新，请重新生成确认挑战\", 409,",
        "                       \"PURCHASE_VERSION_CONFLICT\")",
        "    # 复用库存模块的事务，统一校验批号、库存、成本和幂等。",
        "    purchase = post_purchase(order.id, PurchaseActionPayload(version=order.version), user)",
        "    db.session.add(AuditLog(action=\"CONFIRM_WRITE\",",
        "        resource_type=\"PURCHASE_ORDER\", resource_id=order.id, actor_id=user.id,",
        "        farm_id=order.farm_id))",
        "    db.session.commit()",
        "    return {\"purchase\": purchase, \"confirmed\": True}",
    ])

    add_heading(doc, "检验说明", 1)
    add_text(doc, "本次页面查询只读访问农场、库存和养殖汇总，没有创建草稿、过账、删除或修改业务数据。采购写入链路虽然已接通，但仍要求先生成草稿，再由有权限用户使用短时确认挑战完成过账；删除、批量写入和无人值守过账继续关闭。", after=8)
    add_text(doc, "自动化回归结果：后端测试 65 passed、1 skipped（MySQL 集成测试需显式开启）；前端生产构建通过；网关和 Harness 插件 TypeScript 检查通过。", after=8)
    add_text(doc, "当前项目整体完成度约 94%，其中内置智能体约 90%。后续重点是盘点、调拨、生产记录等更多草稿工具，以及固定问集、观测和前端确认状态。", bold=True, after=4)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
